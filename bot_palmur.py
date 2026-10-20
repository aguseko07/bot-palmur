import telebot
import openpyxl
import os

# Gunakan Token Bot khusus Palmur dari @BotFather
TOKEN = 'KODE_TOKEN_BOT_PALMUR_DARI_BOTFATHER'
bot = telebot.TeleBot(TOKEN)

EXCEL_PALMUR = 'Laporan_Palmur.xlsx'

# Buat file excel template jika belum ada
def init_excel():
    if not os.path.exists(EXCEL_PALMUR):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Palmur"
        # Header Tabel
        headers = ["No", "Nama Item", "Target", "Sales", "Ach (%)"]
        ws.append(headers)
        wb.save(EXCEL_PALMUR)

init_excel()

def format_rp(angka):
    if angka is None or angka == "":
        return "-"
    if isinstance(angka, (int, float)):
        return f"Rp {angka:,.0f}".replace(',', '.')
    return str(angka)

def format_val(val):
    if val is None or val == "":
        return "-"
    if isinstance(val, float):
        return f"{val:.1f}%"
    return str(val)

@bot.message_handler(commands=['start'])
def send_welcome(message):
    bantuan = (
        "🌴 *Bot Laporan Palmur V1.0*\n\n"
        "Perintah yang tersedia:\n"
        "📝 `/palmur <nama_item> <target> <sales>`\n"
        "   *(Contoh: `/palmur MinyakGoreng 500000 450000`)*\n"
        "📋 `/laporan_palmur` - Lihat rekap seluruh item Palmur\n"
        "🗑️ `/hapus_palmur` - Hapus data item terakhir"
    )
    bot.reply_to(message, bantuan, parse_mode='Markdown')

# ==========================================
# COMMAND: /palmur <nama> <target> <sales>
# ==========================================
@bot.message_handler(commands=['palmur'])
def handle_input_palmur(message):
    try:
        args = message.text.split(maxsplit=3)
        if len(args) < 4:
            bot.reply_to(message, "❌ Format salah!\nGunakan: `/palmur <nama_item> <target> <sales>`\nContoh: `/palmur MinyakGoreng 500000 450000`", parse_mode='Markdown')
            return
            
        nama_item = args[1]
        target = float(args[2].replace('.', '').replace(',', ''))
        sales = float(args[3].replace('.', '').replace(',', ''))
        
        # Hitung Achievement (Ach %)
        ach = (sales / target * 100) if target > 0 else 0
        
        wb = openpyxl.load_workbook(EXCEL_PALMUR)
        ws = wb.active
        
        # Tentukan nomor urut berikutnya
        next_row = ws.max_row + 1
        no = next_row - 1
        
        # Masukkan ke Excel
        ws.append([no, nama_item, target, sales, ach])
        wb.save(EXCEL_PALMUR)
        
        pesan_sukses = (
            f"✅ *Item Palmur Berhasil Disimpan!*\n\n"
            f"🏷️ Item: *{nama_item}*\n"
            f"🎯 Target: {format_rp(target)}\n"
            f"💰 Sales: {format_rp(sales)}\n"
            f"📈 Ach: {format_val(ach)}"
        )
        bot.reply_to(message, pesan_sukses, parse_mode='Markdown')
        
    except ValueError:
        bot.reply_to(message, "❌ Pastikan Target dan Sales berupa angka yang valid.")
    except Exception as e:
        bot.reply_to(message, f"❌ Terjadi kesalahan: {e}")

# ==========================================
# COMMAND: /laporan_palmur
# ==========================================
@bot.message_handler(commands=['laporan_palmur'])
def handle_laporan_palmur(message):
    try:
        wb = openpyxl.load_workbook(EXCEL_PALMUR, data_only=True)
        ws = wb.active
        
        if ws.max_row <= 1:
            bot.reply_to(message, "❌ Belum ada data item Palmur yang tersimpan.")
            return

        pesan = "📋 *REKAP LAPORAN PALMUR*\n"
        pesan += "--------------------------------------\n"
        
        total_target = 0
        total_sales = 0
        
        for r in range(2, ws.max_row + 1):
            no = ws.cell(row=r, column=1).value
            item = ws.cell(row=r, column=2).value
            tgt = ws.cell(row=r, column=3).value or 0
            sls = ws.cell(row=r, column=4).value or 0
            ach = ws.cell(row=r, column=5).value or 0
            
            total_target += tgt
            total_sales += sls
            
            pesan += f"{no}. *{item}*\n"
            pesan += f"   Tgt: {format_rp(tgt)} | Sls: {format_rp(sls)} | Ach: *{format_val(ach)}*\n"
            
        total_ach = (total_sales / total_target * 100) if total_target > 0 else 0
        
        pesan += "--------------------------------------\n"
        pesan += f"🎯 *Total Target:* {format_rp(total_target)}\n"
        pesan += f"💵 *Total Sales:* {format_rp(total_sales)}\n"
        pesan += f"📈 *Ach Keseluruhan:* *{format_val(total_ach)}*"
        
        bot.reply_to(message, pesan, parse_mode='Markdown')
        
        with open(EXCEL_PALMUR, 'rb') as doc:
            bot.send_document(message.chat.id, doc, caption="📂 File Excel Palmur Terbaru")
            
    except Exception as e:
        bot.reply_to(message, f"❌ Gagal mengambil laporan: {e}")

# ==========================================
# COMMAND: /hapus_palmur
# ==========================================
@bot.message_handler(commands=['hapus_palmur'])
def handle_hapus_palmur(message):
    try:
        wb = openpyxl.load_workbook(EXCEL_PALMUR)
        ws = wb.active
        
        if ws.max_row <= 1:
            bot.reply_to(message, "❌ Tidak ada data yang bisa dihapus.")
            return
            
        deleted_item = ws.cell(row=ws.max_row, column=2).value
        ws.delete_rows(ws.max_row)
        wb.save(EXCEL_PALMUR)
        
        bot.reply_to(message, f"🗑️ Data item terakhir (*{deleted_item}*) berhasil dihapus.")
    except Exception as e:
        bot.reply_to(message, f"❌ Gagal menghapus: {e}")

if __name__ == '__main__':
    print("🤖 Bot Palmur V1.0 Berhasil Dijalankan...")
    bot.polling(none_stop=True, interval=0)